"""
Build the synthetic billionaire panel and the reference results for module 7.

Module 7 is about the Python around one real SQL step: 01_rtb_ca.sql in
github.com/fhoces/opa-prop40 (bsz-analysis/), which turns daily Forbes
real-time billionaire snapshots into the California lists of a published
paper. The real snapshots are confidential, so this script invents a panel in
the same shape. Every person, id and number here is made up.

This is a Python port of sql-industry-prep/data/build_rtb_sample.R (same seed,
same schema, same hand-placed cases and the same invented ids for them). The
courses do not share files, and numpy's random numbers differ from R's, so
the remaining people and every worth differ from the SQL course's panel.

Writes, from the repo root:
  data/rtb_sample.csv                 rtb_all_combined: one row per (date,
                                      forbes_id), worth in $ million; missing
                                      values are EMPTY fields, and some source
                                      values carry a trailing space, as in the
                                      real bundle's CSVs
  data/rtb_ca_cik.csv                 forbes_id, cik
  data/rtb_residency_overrides.csv    forbes_id, rule, note
  data/expected/                      what module-07/check.py compares with:
    load_profile.csv                    per column: type, NULL count, distinct
                                        values, sum (drill 1)
    exports/<table>.csv                 the three deterministic exports (drill 2)
    answer_key_aggregate.csv            a pretend "authors' sheet" for drill 3
    compare_result.json                 what compare() should return on it
    parity_r/, parity_bad/              two more export folders for drill 4
    parity_results.json                 what parity() should return on them

The reference results come from module-07/solution.py.

Run from the repo root with: python data/build_rtb_sample.py
"""
import csv
import json
import shutil
import sqlite3
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
EXPECTED = DATA / "expected"
sys.path.insert(0, str(ROOT / "module-07"))
import solution  # noqa: E402

rng = np.random.default_rng(2026)

# --- dates: every 8th day, plus the snapshot dates, no 2025-12-31 ---------------
EOY = ["2019-12-31", "2020-12-31", "2021-12-31", "2022-12-31", "2023-12-31",
       "2024-12-31", "2026-01-01"]
SKIP = ["2022-07-18", "2026-03-29", "2026-03-30"]
d0, d1 = date(2019, 12, 31), date(2026, 3, 30)
dates = {(d0 + timedelta(days=k)).isoformat() for k in range(0, (d1 - d0).days + 1, 8)}
dates |= set(EOY) | set(SKIP) | {"2025-12-30"}
dates.discard("2025-12-31")
dates = sorted(dates)
n_dates = len(dates)

# --- people: 60 invented billionaires -------------------------------------------
first = ["Avery", "Bram", "Calla", "Dorian", "Ember", "Faro", "Gideon", "Halle",
         "Ines", "Jory", "Kestrel", "Lumen", "Marlo", "Nell", "Orrin", "Pia",
         "Quill", "Rhea", "Soren", "Tamsin", "Ulla", "Vance", "Wren", "Xavi",
         "Yara", "Zeno", "Adair", "Bexley", "Cyrus", "Delphine"]
last = ["Ashgrove", "Brightwater", "Coldharbor", "Dunmore", "Elderfield",
        "Fairweather", "Glenrock", "Hollowell", "Ironwood", "Juniper",
        "Kettleby", "Larkspur", "Merriweather", "Northcott", "Oakhurst",
        "Pennyroyal", "Quarrington", "Ravensworth", "Silverbrook", "Thistlewood",
        "Underhill", "Vantongeren", "Whitlock", "Yarrowby", "Zanderfell",
        "Amberly", "Blackthorn", "Copperfield", "Driftwood", "Emberton"]
# The hand-placed people keep the SQL course's invented names, so the query's
# top-4 ids and the override table are the same in both courses.
fixed = ["Marlo Zanderfell", "Calla Vantongeren", "Xavi Juniper", "Quill Silverbrook",
         "Wren Underhill", "Soren Driftwood", "Rhea Dunmore"]
pool = [f"{a} {b}" for a in first for b in last if f"{a} {b}" not in fixed]
names = fixed + list(rng.choice(pool, size=60 - len(fixed), replace=False))
n_people = len(names)

industry_sources = [
    ("Technology", "software"), ("Technology", "semiconductors"),
    ("Technology", "online marketplace"), ("Finance & Investments", "hedge funds"),
    ("Finance & Investments", "private equity"), ("Media & Entertainment", "streaming"),
    ("Real Estate", "real estate"), ("Healthcare", "medical devices"),
    ("Food & Beverage", "coffee chains"), ("Fashion & Retail", "athletic apparel"),
    ("Diversified", "investments"),
]
w = np.array([4, 3, 2, 3, 2, 2, 2, 2, 1, 1, 1], dtype=float)
ind = rng.choice(len(industry_sources), size=n_people, p=w / w.sum())
states = rng.choice(["California", "New York", "Texas", "Washington", "Florida",
                     "Nevada", "Massachusetts"], size=n_people,
                    p=[0.50, 0.12, 0.10, 0.08, 0.08, 0.06, 0.06])
people = pd.DataFrame({
    "forbes_name": names,
    "forbes_id": [n.lower().replace(" ", "-") for n in names],
    "state": states.astype(object),
    "country_citizenship": "United States",
    "industries": [industry_sources[i][0] for i in ind],
    "source": [industry_sources[i][1] for i in ind],
})
# Hand-placed cases (rows are 0-based here):
# 0-3 the top 4; 4 California by override (state Nevada); 5 California by
# override (no state); 6 state California but excluded; 7-8 no state, no
# override; 9-10 not US citizens; 11 the Sports owner, never a public split.
people.loc[0:3, "state"] = "California"
people.loc[0:3, "industries"] = "Technology"
people.loc[0:3, "source"] = ["software", "semiconductors", "online marketplace", "software"]
people.loc[4, "state"] = "Nevada"
people.loc[5, "state"] = None
people.loc[6, "state"] = "California"
people.loc[7:8, "state"] = None
people.loc[9:10, "state"] = None
people.loc[9:10, "country_citizenship"] = ["Canada", "Brazil"]
people.loc[11, ["state", "industries", "source"]] = ["California", "Sports", "sports team"]

w0 = np.exp(rng.uniform(np.log(600), np.log(40000), n_people))
w0[0:4] = [160000, 120000, 95000, 90000]
w0[[4, 5, 6, 11]] = [4000, 2500, 6000, 1800]
public_share = np.round(rng.uniform(0, 0.95, n_people), 3)
first_day = np.zeros(n_people, dtype=int)
last_day = np.full(n_people, n_dates - 1)
late = rng.choice(np.arange(12, n_people), size=10, replace=False)
first_day[late] = rng.integers(29, 250, size=10)
gone = rng.choice(np.setdiff1d(np.arange(12, n_people), late), size=2, replace=False)
last_day[gone] = rng.integers(149, 280, size=2)

# --- the panel: a random walk in log worth per person ---------------------------
rows = []
for i, p in people.iterrows():
    days = np.arange(first_day[i], last_day[i] + 1)
    steps = rng.normal(0.003, 0.04, size=len(days))
    worth = np.round(w0[i] * np.exp(np.cumsum(steps) - steps[0]), 1)
    for d, wv in zip(days, worth):
        pub = round(wv * public_share[i], 1)
        rows.append([dates[d], p.forbes_id, p.forbes_name, p.state, p.country_citizenship,
                     p.source, p.industries, wv, pub, round(wv - pub, 1)])
cols = ["date", "forbes_id", "forbes_name", "state", "country_citizenship", "source",
        "industries", "forbes_worth", "forbes_public_worth", "forbes_private_worth"]
panel = pd.DataFrame(rows, columns=cols).astype(object)
sports = panel["forbes_id"] == people.loc[11, "forbes_id"]
no_split = rng.uniform(size=len(panel)) < 0.03
panel.loc[sports | no_split, ["forbes_public_worth", "forbes_private_worth"]] = None
panel.loc[rng.uniform(size=len(panel)) < 0.005, "forbes_worth"] = None
panel = panel.sort_values(["date", "forbes_id"]).reset_index(drop=True)
# The real CSVs carry a trailing space on many source values; so does this one.
trail = rng.uniform(size=len(panel)) < 0.2
panel.loc[trail, "source"] = panel.loc[trail, "source"] + " "

overrides = pd.DataFrame({
    "forbes_id": people.loc[[4, 5, 6], "forbes_id"].to_list(),
    "rule": ["include", "include", "exclude"],
    "note": ["course example: state on file is Nevada, counted as California",
             "course example: no state on file, counted as California",
             "course example: state on file is California, never counted"],
})
ca_ids = people.loc[(people["state"] == "California")
                    | people["forbes_id"].isin(overrides.loc[overrides.rule == "include",
                                                             "forbes_id"]), "forbes_id"]
cik = pd.DataFrame({"forbes_id": ca_ids.to_list(),
                    "cik": rng.choice(np.arange(1000000, 2000000), size=len(ca_ids),
                                      replace=False).astype(object)})
cik.loc[rng.uniform(size=len(cik)) < 0.25, "cik"] = None


def write(df, path):
    # None -> empty field, as the real bundle writes missing values.
    df.to_csv(path, index=False, na_rep="", lineterminator="\n")


write(panel, DATA / "rtb_sample.csv")
write(cik, DATA / "rtb_ca_cik.csv")
write(overrides, DATA / "rtb_residency_overrides.csv")
print(f"data/rtb_sample.csv: {len(panel):,} rows ({n_people} people x {n_dates} dates); "
      f"rtb_ca_cik {len(cik)}, overrides {len(overrides)}")

# --- reference results -------------------------------------------------------------
RTB_COLUMNS = [("date", "TEXT"), ("forbes_id", "TEXT"), ("forbes_name", "TEXT"),
               ("state", "TEXT"), ("country_citizenship", "TEXT"), ("source", "TEXT"),
               ("industries", "TEXT"), ("forbes_worth", "REAL"),
               ("forbes_public_worth", "REAL"), ("forbes_private_worth", "REAL")]
EXPORT_ORDER = {
    "rtb_ca_eoy": "date, forbes_worth DESC, forbes_id",
    "rtb_ca_2026_01_01_industry": "(industries = 'Total'), fraction_forbes_worth DESC, industries",
    "rtb_ca_aggregate": "date",
}
if EXPECTED.exists():
    shutil.rmtree(EXPECTED)
EXPECTED.mkdir(parents=True)

with tempfile.TemporaryDirectory() as tmp:
    con = sqlite3.connect(Path(tmp) / "rtb.sqlite")
    n = solution.load_csv(con, DATA / "rtb_sample.csv", "rtb_all_combined", RTB_COLUMNS,
                          index_on=("date", "forbes_id"))
    solution.load_csv(con, DATA / "rtb_ca_cik.csv", "rtb_ca_cik",
                      [("forbes_id", "TEXT"), ("cik", "INTEGER")])
    solution.load_csv(con, DATA / "rtb_residency_overrides.csv", "rtb_residency_overrides",
                      [("forbes_id", "TEXT"), ("rule", "TEXT"), ("note", "TEXT")])

    # Drill 1: a profile of the loaded table.
    prof = []
    for name, typ in RTB_COLUMNS:
        n_null, n_distinct = con.execute(
            f"SELECT SUM({name} IS NULL), COUNT(DISTINCT {name}) FROM rtb_all_combined").fetchone()
        total = con.execute(f"SELECT SUM({name}) FROM rtb_all_combined").fetchone()[0] \
            if typ == "REAL" else ""
        prof.append({"column": name, "type": typ, "n_rows": n, "n_null": n_null,
                     "n_distinct": n_distinct, "sum": total})
    pd.DataFrame(prof).to_csv(EXPECTED / "load_profile.csv", index=False, lineterminator="\n")

    # Drill 2: the three exports.
    counts = solution.run_and_export(con, ROOT / "module-07" / "01_rtb_ca.sql",
                                     list(EXPORT_ORDER), EXPECTED / "exports", EXPORT_ORDER)
    con.close()
for t, k in counts.items():
    print(f"  data/expected/exports/{t}.csv: {k} rows")

# Drill 3: a pretend authors' sheet for the daily aggregate, with a vintage gap
# (it stops 5 dates earlier), one cell off by 1e-9 (noise, within tolerance) and
# one cell off by 0.25 (a real difference).
agg = pd.read_csv(EXPECTED / "exports" / "rtb_ca_aggregate.csv")
key = agg.iloc[:-5].copy()
key = key.drop(columns=["forbes_private_worth"])  # the sheet does not have it
key.loc[10, "forbes_worth_total"] += 1e-9
key.loc[40, "forbes_worth_top4"] += 0.25
key.to_csv(EXPECTED / "answer_key_aggregate.csv", index=False, float_format="%.17g",
           lineterminator="\n")
res = solution.compare(agg, key, ["date"], ["n_billionaires", "forbes_worth_total",
                                            "forbes_worth_top4"])
(EXPECTED / "compare_result.json").write_text(json.dumps(res, indent=1) + "\n")
print("  compare_result:", res)


# Drill 4: an "R" export folder (same doubles, written with 17 significant
# digits, as R's writer sometimes does) and a "bad" one (one value moved by
# 1e-6, beyond the 1e-9 parity tolerance).
def rewrite(src, dst, change=None):
    dst.mkdir(parents=True, exist_ok=True)
    for t in EXPORT_ORDER:
        with open(src / f"{t}.csv", newline="") as f:
            rows = list(csv.reader(f))
        for r in rows[1:]:
            for j, v in enumerate(r):
                try:
                    x = float(v)
                except ValueError:
                    continue
                if "." in v or "e" in v:
                    r[j] = "%.17g" % x
        if change and t == change[0]:
            r = rows[change[1]]
            r[change[2]] = repr(float(r[change[2]]) + 1e-6)
        with open(dst / f"{t}.csv", "w", newline="") as f:
            csv.writer(f, lineterminator="\n").writerows(rows)


rewrite(EXPECTED / "exports", EXPECTED / "parity_r")
rewrite(EXPECTED / "exports", EXPECTED / "parity_bad", change=("rtb_ca_aggregate", 5, 2))
par = {
    "parity_r": solution.parity(EXPECTED / "exports", EXPECTED / "parity_r", list(EXPORT_ORDER)),
    "parity_bad": solution.parity(EXPECTED / "exports", EXPECTED / "parity_bad", list(EXPORT_ORDER)),
}
(EXPECTED / "parity_results.json").write_text(json.dumps(par, indent=1) + "\n")
print("  parity_results:", {k: {t: v["ok"] for t, v in d.items()} for k, d in par.items()})


# =================================================================================
# Part 2 (drills 5 to 8): a small workbook, a yearly panel, two export folders
# =================================================================================
# A second generator, so the part 1 files above stay byte-identical when this
# section changes.
import datetime  # noqa: E402
import math  # noqa: E402
import re  # noqa: E402
import zipfile  # noqa: E402

import openpyxl  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

rng2 = np.random.default_rng(4026)
WORKBOOK = DATA / "workbook_sample.xlsx"
MONEY = ["forbes_worth", "forbes_public_worth", "purchase", "sale", "kg", "kg_long",
         "kg_short", "option_profit", "noneq_comp", "ordinary_income", "kg_taxable",
         "dividend", "fiscal_income", "donation", "donation_deductible", "income_taxable",
         "ca_income_tax", "fed_ordinary_income_tax", "fed_preferential_tax",
         "fed_income_tax", "fiscal_income_tax", "sales_tax", "w_txt", "w_tax_ppent",
         "w_pi", "public_worth", "public_worth_avg", "total_tax", "economic_income"]
SUMMED = [c for c in MONEY if c not in ("public_worth", "public_worth_avg")]   # 27

# data_sec_all: one row per California billionaire and year, $ million, in
# sheet order, with the cases the R code had to handle (see the SQL course's
# build script for the same design): an excluded id, two fortunes for one
# person in 2022, a column empty for all of 2019, and a re-pasted tail of
# four 2025 rows (one with a re-typed dividend, one with no worth).
ca = list(ca_ids)
w_start = np.exp(rng2.uniform(np.log(1100), np.log(30000), len(ca)))
w_start[:4] = [150000, 110000, 90000, 85000]           # the top 4 come first in ca
SHARES = {"purchase": .004, "sale": .010, "kg": .008, "kg_long": .007, "kg_short": .001,
          "option_profit": .002, "noneq_comp": .0005, "ordinary_income": .001,
          "kg_taxable": .008, "dividend": .003, "fiscal_income": .012, "donation": .002,
          "donation_deductible": .001, "income_taxable": .010, "ca_income_tax": .0012,
          "fed_ordinary_income_tax": .0004, "fed_preferential_tax": .0015,
          "fed_income_tax": .0019, "fiscal_income_tax": .003, "sales_tax": .0002,
          "w_txt": .004, "w_tax_ppent": .0003, "w_pi": .03, "total_tax": .008,
          "economic_income": .04}
recs = []
for yr in range(2019, 2026):
    keep = (rng2.uniform(size=len(ca)) < 0.85) | (np.arange(len(ca)) < 4)
    for i in np.flatnonzero(keep):
        wv = round(float(w_start[i] * np.exp(rng2.normal(0.06 * (yr - 2019), 0.15))), 3)
        pub = round(wv * float(rng2.uniform(0.3, 0.95)), 3)
        r = {"year": yr, "forbes_id": ca[i], "forbes_worth": wv, "forbes_public_worth": pub}
        for c, sh in SHARES.items():
            v = round(wv * sh * float(rng2.uniform(0.2, 1.8)), 4)
            p_null = 0 if c in ("w_txt", "w_tax_ppent", "w_pi", "total_tax", "economic_income") \
                else 0.5 if c == "option_profit" else 0.15
            r[c] = None if rng2.uniform() < p_null else v
        r["public_worth"] = round(pub * float(rng2.uniform(0.95, 1.05)), 3)
        r["public_worth_avg"] = round(pub * float(rng2.uniform(0.85, 1.0)), 3)
        recs.append(r)
dsa = pd.DataFrame(recs)[["year", "forbes_id"] + MONEY]
others = [i for i in ca if i not in set(people.loc[0:11, "forbes_id"])]
excluded_id, twin_id = others[0], others[1]
dsa = dsa[~((dsa.forbes_id == twin_id) & (dsa.year == 2022))]
twin = pd.DataFrame({"year": 2022, "forbes_id": twin_id, "forbes_worth": [8000.0, 5256.0]})
twin["forbes_public_worth"] = (twin["forbes_worth"] * 0.6).round(3)
for c in MONEY[2:]:
    twin[c] = (twin["forbes_worth"] * rng2.uniform(0.001, 0.01, 2)).round(4)
dsa = pd.concat([dsa, twin], ignore_index=True)
dsa = dsa.sort_values(["year", "forbes_worth"], ascending=[True, False], kind="stable")
dsa = dsa.reset_index(drop=True).astype(object)
dsa.loc[dsa.year == 2019, "option_profit"] = None
last4 = list(dsa.index[(dsa.year == 2025) & (dsa.forbes_id != excluded_id)][-4:])
dsa.loc[last4[3], "forbes_worth"] = None
tail = dsa.loc[last4].copy()
tail.loc[last4[1], "dividend"] = round(float(tail.loc[last4[1], "dividend"]) + 0.5, 4)
dsa = pd.concat([dsa, tail], ignore_index=True)

# The workbook. Sheet data_sec_all: three title rows, the header in row 4,
# data from row 5, an empty cell for every missing value. Sheet summary: a
# hand-made block of labels and numbers, some typed as text, as positional
# sheets in the real workbook are.
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "data_sec_all"
ws["A1"] = "Back to index"
ws["A2"] = "Wealth, income and taxes of California billionaires (synthetic, course data)"
ws["A3"] = "All dollar value variables in $ million (nominal)"
ws.append(["year", "forbes_id"] + MONEY)
for rec in dsa.itertuples(index=False):
    ws.append([None if (v is None or (isinstance(v, float) and math.isnan(v))) else v
               for v in rec])
sm = wb.create_sheet("summary")
sm["A1"] = "Back to index"
sm["A2"] = "Income tax block (synthetic, course data)"
sm["AB2"] = "note: this sheet is 28 columns wide"
sm.append([])
sm["A4"] = "year"
for j, yr in enumerate(range(2018, 2023)):
    sm.cell(row=4, column=2 + j, value=yr)
    sm.cell(row=5, column=2 + j, value=round(1800 + 95.5 * j + float(rng2.uniform(0, 30)), 3))
sm["A5"] = "ca_agi_b"
sm["A6"] = "ca_inctax_b"
for col, v in zip("BCDEF", [" 97.293 ", "1_000", "n/a", None, "1e3"]):
    sm[f"{col}6"] = v
sm["A7"] = "as_of"
sm["B7"] = datetime.datetime(2018, 12, 31)
sm["C7"] = True
sm["D7"] = "  -12  "
sm["E7"] = "1,234"
sm["F7"] = 0.1 + 0.2
sm["A12"] = "Total"
sm["B12"] = 12345.678
sm["A25"].font = Font(bold=True)   # a formatted but empty cell: a trailing empty row
for w_ in (wb.properties,):
    w_.creator = "build_rtb_sample.py"
    w_.created = w_.modified = datetime.datetime(2026, 1, 1)
tmp_xlsx = DATA / "workbook_sample.tmp.xlsx"
wb.save(tmp_xlsx)
# Rewrite the zip with fixed timestamps, so a rebuild gives identical bytes.
with zipfile.ZipFile(tmp_xlsx) as zin, zipfile.ZipFile(WORKBOOK, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        info = zipfile.ZipInfo(item.filename, date_time=(2026, 1, 1, 0, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        data = zin.read(item.filename)
        if item.filename == "docProps/core.xml":   # openpyxl stamps the save time here
            data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
        zout.writestr(info, data)
tmp_xlsx.unlink()
print(f"data/workbook_sample.xlsx: data_sec_all {len(dsa)} rows (excluded id {excluded_id}), "
      "summary sheet")

# Drill 5: what read_sheet() should return, checked against what was written.
summary = solution.read_sheet(WORKBOOK, "summary")
assert summary.shape == (12, 28), summary.shape
assert summary.at[7, "B"] == 43465.0 and summary.at[7, "C"] is True
sheet_checks = {
    "summary": {"shape": list(summary.shape), "columns_last": summary.columns[-1],
                "cells": {"A2": summary.at[2, "A"], "AB2": summary.at[2, "AB"],
                          "B4": summary.at[4, "B"], "B6": summary.at[6, "B"],
                          "E6": summary.at[6, "E"], "B7": summary.at[7, "B"],
                          "C7": summary.at[7, "C"], "B12": summary.at[12, "B"]}},
}
dsa_sheet = solution.read_sheet(WORKBOOK, "data_sec_all")
sheet_checks["data_sec_all"] = {
    "shape": list(dsa_sheet.shape), "columns_last": dsa_sheet.columns[-1],
    "cells": {"A4": dsa_sheet.at[4, "A"], "AE4": dsa_sheet.at[4, "AE"],
              "A5": dsa_sheet.at[5, "A"], "B5": dsa_sheet.at[5, "B"],
              f"C{len(dsa_sheet)}": dsa_sheet.at[len(dsa_sheet), "C"]},
}
(EXPECTED / "read_sheet.json").write_text(json.dumps(sheet_checks, indent=1) + "\n")

# Drill 6: what xls_cell() should return for a list of addresses.
ADDRS = ["B4", "B5", "F5", "A5", "B6", "C6", "D6", "E6", "F6", "B7", "C7", "D7", "E7",
         "F7", "B12", "AB2", "Z99", "A30"]
cells = [[a, None if math.isnan(x) else x] for a, x in
         ((a, solution.xls_cell(summary, a)) for a in ADDRS)]
assert dict(cells)["B6"] == 97.293 and dict(cells)["C6"] is None and dict(cells)["F6"] == 1000.0
(EXPECTED / "xls_cells.json").write_text(json.dumps(cells, indent=1) + "\n")

# Drill 7: the yearly aggregates from the real SQL file, the reference that
# the pandas twin must match. The input is the sheet as pandas reads it.
dsa_in = pd.read_excel(WORKBOOK, sheet_name="data_sec_all", skiprows=3)
inputs7 = {"exclude_ids": [excluded_id], "cols": SUMMED}
(EXPECTED / "data_sec_agg_inputs.json").write_text(json.dumps(inputs7, indent=1) + "\n")
with tempfile.TemporaryDirectory() as tmp:
    con = sqlite3.connect(Path(tmp) / "wb.sqlite")
    d = dsa_in.copy()
    d.insert(0, "row_num", np.arange(1, len(d) + 1))
    d.to_sql("data_sec_all", con, index=False)
    pd.DataFrame({"forbes_id": [excluded_id]}).to_sql("data_sec_agg_exclude", con, index=False)
    con.executescript((ROOT / "module-07" / "02_data_sec_agg.sql").read_text(encoding="utf-8"))
    agg = pd.read_sql("SELECT * FROM data_sec_agg ORDER BY year", con)
    con.close()
write_csv = lambda df, path: df.to_csv(path, index=False, lineterminator="\n",  # noqa: E731
                                       float_format=None, na_rep="NA")
write_csv(agg, EXPECTED / "data_sec_agg.csv")
twin_check = solution.compute_data_sec_agg(dsa_in, [excluded_id], SUMMED)
rel = (twin_check[SUMMED].to_numpy(float) - agg[SUMMED].to_numpy(float))
assert np.all(np.abs(rel) <= 1e-9 * np.maximum(1, np.abs(agg[SUMMED].to_numpy(float))))
assert list(twin_check["n"]) == list(agg["n"])
print(f"  data/expected/data_sec_agg.csv: {len(agg)} years, from 02_data_sec_agg.sql")


# Drill 8: an "R" export folder and two "Python" ones. R writes 17 significant
# digits and NA; the good Python folder writes repr() and differs only by
# rounding noise; the bad one has four real problems.
def write_rows(path, header, rows, fmt):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow([fmt(v) for v in r])


def fmt_r(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return "NA"
    return "%.17g" % v if isinstance(v, float) else str(v)


def fmt_py(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return "NA"
    return repr(v) if isinstance(v, float) else str(v)


PAR = EXPECTED / "parity_rel"
agg_rows = [[int(r[0]), int(r[1])] + [float(x) for x in r[2:]]
            for r in agg.itertuples(index=False)]
ftb_rows = [[yr, int(rng2.integers(17_000_000, 21_000_000)),
             float(rng2.uniform(1.8e12, 2.3e12)).__round__(2)] for yr in range(2016, 2023)]
rates_rows = []
for period in ("2004-2016", "2017-2025"):
    tot = float(rng2.uniform(0.2, 0.3))
    parts = [float(x) for x in rng2.uniform(0.02, 0.06, 4)]
    rest = tot - sum(parts)
    check = tot - (parts[0] + parts[1] + parts[2] + parts[3] + rest)   # about 1e-17
    rates_rows.append([period, tot, rest, check + 1.3e-13, None if period == "2004-2016"
                       else float(rng2.uniform(40000, 60000))])
files = {
    "data_sec_agg.csv": (list(agg.columns), agg_rows),
    "ftb_totals.csv": (["taxable_year", "all_returns", "ca_agi"], ftb_rows),
    "tax_rates.csv": (["period", "total_tax_per_income", "other_per_income",
                       "check_income_decomp", "avg_wealth_m"], rates_rows),
}
if PAR.exists():
    shutil.rmtree(PAR)
for name, (header, rows) in files.items():
    write_rows(PAR / "r" / name, header, rows, fmt_r)

good = {k: (h, [list(r) for r in rows]) for k, (h, rows) in files.items()}
good["data_sec_agg.csv"][1][2][4] *= 1 + 2e-15          # summation-order noise
good["ftb_totals.csv"][1][3][2] = float(np.nextafter(np.nextafter(
    good["ftb_totals.csv"][1][3][2], np.inf), np.inf))    # 2 ulps: ~5e-4 apart in $
good["tax_rates.csv"][1][1][3] += 3e-16                 # tiny value, tiny absolute gap
for name, (header, rows) in good.items():
    write_rows(PAR / "py_good" / name, header, rows, fmt_py)

bad = {k: (h, [list(r) for r in rows]) for k, (h, rows) in files.items()}
bad["data_sec_agg.csv"][1][5][6] *= 1 + 1e-7            # a real difference
del bad["ftb_totals.csv"]                               # a file Python never wrote
bad["tax_rates.csv"][1][0][4] = 0.0                     # NA in R, 0 in Python
bad["notes.csv"] = (["note"], [["written by Python only"]])
for name, (header, rows) in bad.items():
    write_rows(PAR / "py_bad" / name, header, rows, fmt_py)

par2 = {k: solution.parity_rel(PAR / "r", PAR / k) for k in ("py_good", "py_bad")}
assert all(v["ok"] for v in par2["py_good"].values()), par2["py_good"]
assert not any(v["ok"] for v in par2["py_bad"].values()), par2["py_bad"]
(EXPECTED / "parity_rel_results.json").write_text(json.dumps(par2, indent=1) + "\n")
print("  parity_rel_results:", {k: {t: v["ok"] for t, v in d.items()} for k, d in par2.items()})
